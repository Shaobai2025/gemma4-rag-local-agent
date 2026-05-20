from io import BytesIO
from docx import Document
import pandas as pd

def markdown_bytes(title: str, body: str) -> bytes:
    return f"# {title}\n\n{body}\n".encode("utf-8")

def docx_bytes(title: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        doc.add_paragraph(para)
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()

def excel_report_bytes(answer: str, preview_markdown: str, selected_sheet: str, analyzed_columns: list[str]) -> bytes:
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        pd.DataFrame([{
            "selected_sheet": selected_sheet,
            "analyzed_columns": ", ".join(analyzed_columns),
            "analysis_answer": answer,
        }]).to_excel(writer, index=False, sheet_name="analysis")
        pd.DataFrame({"preview_markdown": [preview_markdown]}).to_excel(writer, index=False, sheet_name="preview")
    bio.seek(0)
    return bio.read()
