from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from fastapi.responses import Response

from app.excel_agent.session_store import load_json, load_text, session_dir


def export_markdown_response(session_id: str) -> Response:
    text = load_text(session_id, "last_report.md")
    return Response(
        content=text.encode("utf-8"),
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="excel_agent_report.md"'},
    )


def export_word_response(session_id: str) -> Response:
    text = load_text(session_id, "last_report.md")

    try:
        from app.exporters import docx_bytes
        data = docx_bytes("Excel智能分析报告", text)
        return Response(
            content=data,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="excel_agent_report.docx"'},
        )
    except Exception:
        # 简易 docx 兜底
        from docx import Document
        doc = Document()
        doc.add_heading("Excel智能分析报告", level=1)
        for line in text.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            else:
                doc.add_paragraph(line)
        bio = io.BytesIO()
        doc.save(bio)
        return Response(
            content=bio.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="excel_agent_report.docx"'},
        )


def export_pdf_response(session_id: str) -> Response:
    text = load_text(session_id, "last_report.md")
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font_name = "STSong-Light"
        except Exception:
            font_name = "Helvetica"

        c.setFont(font_name, 11)
        x, y = 50, height - 50
        line_height = 16

        for raw_line in text.splitlines():
            line = raw_line.replace("#", "").strip()
            if not line:
                y -= line_height
                continue

            chunks = _wrap_line(line, 42)
            for chunk in chunks:
                if y < 50:
                    c.showPage()
                    c.setFont(font_name, 11)
                    y = height - 50
                c.drawString(x, y, chunk)
                y -= line_height

        c.save()
        return Response(
            content=buffer.getvalue(),
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="excel_agent_report.pdf"'},
        )
    except Exception as e:
        return Response(
            content=f"PDF导出失败：{e}\n请先安装 reportlab：pip install reportlab".encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            status_code=500,
        )


def export_excel_response(session_id: str) -> Response:
    try:
        data = load_json(session_id, "last_artifacts.json")
    except Exception:
        data = {}

    tables = data.get("tables", {})
    warning = tables.get("warning_students", [])

    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        if warning:
            pd.DataFrame(warning).to_excel(writer, sheet_name="预警学生名单", index=False)
        else:
            pd.DataFrame([{"提示": "暂无可导出的结果表，请先执行成绩预警分析。"}]).to_excel(writer, sheet_name="结果", index=False)

    return Response(
        content=bio.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="excel_agent_results.xlsx"'},
    )


def _wrap_line(line: str, n: int) -> List[str]:
    return [line[i:i+n] for i in range(0, len(line), n)] or [""]
